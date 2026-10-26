"""Arma respuestas-ejemplo.xlsx: la planilla que genera el formulario, con datos INVENTADOS.

    python armar-ejemplo.py

Tiene las columnas exactas (ver formulario.md) y casos para probar generar.py:
uno que no muestra precios y usa "Otros datos", una respuesta sin aprobar, otra
completada de cero que todavía no se revisó, un WhatsApp mal escrito y alguien
que intenta meter código en una descripción.
"""
import pathlib
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

AQUI = pathlib.Path(__file__).resolve().parent
MAX = 12
TODO = "Precio, Descripción, Foto, Un botón de WhatsApp en cada producto"

TITULOS = ["Marca temporal", "Dirección de correo electrónico", "Nombre del emprendimiento", "Rubro",
           "Contanos qué hacés", "Logo", "WhatsApp para consultas", "Instagram", "Estilo del catálogo",
           "¿Qué querés mostrar de tus productos?"]
for n in range(1, MAX + 1):
    TITULOS += [f"Producto {n} · Nombre", f"Producto {n} · Precio", f"Producto {n} · Descripción",
                f"Producto {n} · Otros datos", f"Producto {n} · Foto"]
    if n < MAX:
        TITULOS.append(f"Producto {n} · ¿Agregás otro?")
TITULOS += ["Nombre y apellido", "¿En qué curso del CFP estás?", "Conformidad",
            "Aprobado", "Revisó", "Observaciones"]

CONFORME = ("Acepto que los datos del emprendimiento se publiquen en internet en un catálogo del CFP 401. "
            "Mi nombre y mi curso no se publican., Las fotos son mías o tengo permiso para usarlas., "
            "Entiendo que alguien del CFP lo revisa antes de publicarlo, y que puedo pedir que lo den de baja cuando quiera.")


def foto(n):
    return f"https://drive.google.com/open?id=EJEMPLO-{n:04d}"


# Cada producto: (nombre, precio, descripción, otros datos)
RESPUESTAS = [
    {"cuando": datetime(2026, 10, 28, 14, 5), "nombre": "Dulces de la Isla", "rubro": "Dulces caseros",
     "lema": "Dulce de leche y mermeladas hechas en la Isla Santiago.", "wsp": "2215550001",
     "ig": "dulcesdelaisla", "estilo": "Dulce", "mostrar": TODO, "curso": "Pastelería", "aprobado": True,
     "reviso": "Profe",
     "productos": [("Dulce de leche casero", "3500", "Frasco de 500 g.", ""),
                   ("Mermelada de frutilla", "3200", "Con frutillas de la quinta.", "Sin TACC: sí"),
                   ("Alfajores de maicena", "", "Por docena. Consultá por cantidades.", "")]},
    {"cuando": datetime(2026, 10, 28, 15, 40), "nombre": "Hilos y Agujas", "rubro": "Tejidos a mano",
     "lema": "", "wsp": "2215550002", "ig": "", "estilo": "Natural", "mostrar": TODO,
     "curso": "Pensamiento Computacional", "aprobado": True, "reviso": "Preceptoría",
     "productos": [("Gorro de lana", "8000", "Varios colores.", "Talles: S, M y L\nMaterial: lana de oveja"),
                   ("Bufanda", "9500", "", "Largo: 1,60 m")]},
    # No muestra precios ni un botón por producto, y cuenta el material y las medidas.
    {"cuando": datetime(2026, 10, 29, 9, 10), "nombre": "Impresiones 3D Ensenada", "rubro": "Impresión 3D",
     "lema": "Piezas a medida y objetos de diseño.", "wsp": "2215550004", "ig": "", "estilo": "Mar",
     "mostrar": "Descripción, Foto", "curso": "Pensamiento Computacional", "aprobado": True, "reviso": "Profe",
     "productos": [("Maceta geométrica", "", "Para suculentas.",
                    "Material: PLA biodegradable\nMedidas: 10 x 10 x 12 cm\nColores: blanco, verde o terracota"),
                   ("Llavero con tu nombre", "", "Con el nombre que quieras.", "Tiempo de entrega: 3 días")]},
    {"cuando": datetime(2026, 10, 29, 10, 15), "nombre": "La Vela Encendida", "rubro": "Velas aromáticas",
     "lema": "Velas de soja con esencias naturales.", "wsp": "2215550003", "ig": "", "estilo": "Mar",
     "mostrar": TODO, "curso": "Auxiliar en Cosmética", "aprobado": False, "reviso": "",
     "observaciones": "Falta revisar las fotos",
     "productos": [("Vela de lavanda", "4200", "Dura 30 horas.", "")]},
    # Completó el formulario DE CERO (en vez de editar): es una respuesta nueva, sin revisar. Sigue la de arriba.
    {"cuando": datetime(2026, 10, 30, 9, 30), "nombre": "Dulces de la Isla", "rubro": "Dulces caseros",
     "lema": "Dulce de leche y mermeladas hechas en la Isla Santiago.", "wsp": "2215550001",
     "ig": "dulcesdelaisla", "estilo": "Dulce", "mostrar": TODO, "curso": "Pastelería", "aprobado": False,
     "reviso": "",
     "productos": [("Dulce de leche casero", "3800", "Frasco de 500 g.", ""),
                   ("Pastafrola", "6000", "De membrillo o de batata.", "")]},
    # Aprobada, pero con el WhatsApp mal (alguien lo editó a mano): generar.py la frena.
    {"cuando": datetime(2026, 10, 30, 11, 0), "nombre": "Mates del Río", "rubro": "Mates y bombillas",
     "lema": "", "wsp": "221555", "ig": "", "estilo": "Elegante", "mostrar": TODO, "curso": "Marroquinería",
     "aprobado": True, "reviso": "Profe", "productos": [("Mate de calabaza", "7800", "Curado y listo para usar.", "")]},
    # Alguien intenta meter código: tiene que salir como texto.
    {"cuando": datetime(2026, 10, 30, 12, 20), "nombre": "Bijou Ensenada", "rubro": "Bijouterie",
     "lema": "<script>alert('hola')</script>", "wsp": "2215550005", "ig": "@bijou.ensenada", "estilo": "Alegre",
     "mostrar": TODO, "curso": "Pensamiento Computacional", "aprobado": True, "reviso": "Profe",
     "productos": [("Aros de acero", "2500", "Hipoalergénicos.", "<b>Oferta</b>: <img src=x onerror=alert(1)>"),
                   ("Pulsera tejida", "1800", "<b>en negrita?</b>", "")]},
]

libro = openpyxl.Workbook()
hoja = libro.active
hoja.title = "Respuestas de formulario 1"
hoja.append(TITULOS)
for celda in hoja[1]:
    celda.font = Font(bold=True)
for i, r in enumerate(RESPUESTAS, start=1):
    fila = {"Marca temporal": r["cuando"], "Dirección de correo electrónico": f"persona{i}@example.com",
            "Nombre del emprendimiento": r["nombre"], "Rubro": r["rubro"], "Contanos qué hacés": r["lema"],
            "Logo": "", "WhatsApp para consultas": r["wsp"], "Instagram": r["ig"],
            "Estilo del catálogo": r["estilo"], "¿Qué querés mostrar de tus productos?": r["mostrar"],
            "Nombre y apellido": f"Persona de ejemplo {i}", "¿En qué curso del CFP estás?": r["curso"],
            "Conformidad": CONFORME, "Aprobado": r["aprobado"], "Revisó": r["reviso"],
            "Observaciones": r.get("observaciones", "")}
    for n, (nombre, precio, descripcion, otros) in enumerate(r["productos"], start=1):
        fila[f"Producto {n} · Nombre"] = nombre
        fila[f"Producto {n} · Precio"] = int(precio) if precio else None
        fila[f"Producto {n} · Descripción"] = descripcion
        fila[f"Producto {n} · Otros datos"] = otros
        fila[f"Producto {n} · Foto"] = foto(i * 100 + n)
        if n < MAX:
            fila[f"Producto {n} · ¿Agregás otro?"] = "Sí" if n < len(r["productos"]) else "No"
    hoja.append([fila.get(t) for t in TITULOS])
    hoja.cell(row=i + 1, column=1).number_format = "dd/mm/yyyy hh:mm:ss"

# Las tres columnas de revisión, a la derecha de todo, pintadas para que se vean.
for col in range(len(TITULOS) - 2, len(TITULOS) + 1):
    for fila in range(1, len(RESPUESTAS) + 2):
        hoja.cell(row=fila, column=col).fill = PatternFill("solid", fgColor="FFF2CC")
hoja.freeze_panes = "D2"
libro.save(AQUI / "respuestas-ejemplo.xlsx")
print(f"respuestas-ejemplo.xlsx: {len(RESPUESTAS)} respuestas, {len(TITULOS)} columnas")
